import pytest
from unittest.mock import patch, MagicMock
from app.services.extraction.image_extractor import (
    extract_images_from_pdf,
    extract_images_from_docx,
    extract_images_from_pptx,
    extract_images_from_file,
)

@patch("app.services.extraction.image_extractor.ocr_image", return_value="ocr text")
@patch("app.services.extraction.image_extractor.Image.open")
def test_extract_images_from_pdf(mock_image_open, mock_ocr_image):
    mock_pil = MagicMock()
    mock_pil.size = (100, 100)
    mock_image_open.return_value = mock_pil

    with patch("fitz.open") as mock_open:
        mock_doc = MagicMock()
        mock_open.return_value = mock_doc
        
        mock_page = MagicMock()
        mock_doc.__len__.return_value = 1
        mock_doc.__getitem__.return_value = mock_page
        
        # get_images returns list of img_info. [xref, smask, width, height, bpc, colorspace, alt. colorspace, name, filter]
        mock_page.get_images.return_value = [[123]]
        
        mock_doc.extract_image.return_value = {
            "image": b"fakeimage",
            "ext": "png",
            "width": 100,
            "height": 100
        }
        
        images = extract_images_from_pdf(file_path="test.pdf")
        assert len(images) == 1
        assert images[0].width == 100
        assert images[0].ocr_text == "ocr text"

        # Test exception in ocr
        mock_ocr_image.side_effect = Exception("ocr fail")
        images2 = extract_images_from_pdf(file_bytes=b"fake")
        assert len(images2) == 1
        assert images2[0].ocr_text == ""
        mock_ocr_image.side_effect = None

@patch("app.services.extraction.image_extractor.ocr_image", return_value="ocr text")
@patch("app.services.extraction.image_extractor.Image.open")
def test_extract_images_from_docx(mock_image_open, mock_ocr_image):
    mock_pil = MagicMock()
    mock_pil.size = (100, 100)
    mock_pil.format = "jpeg"
    mock_image_open.return_value = mock_pil

    with patch("docx.Document") as mock_doc_class:
        mock_doc = MagicMock()
        mock_doc_class.return_value = mock_doc
        
        mock_rel = MagicMock()
        mock_rel.reltype = "image"
        mock_rel.target_part.blob = b"fakeimage"
        
        mock_doc.part.rels.values.return_value = [mock_rel]
        
        images = extract_images_from_docx(file_path="test.docx")
        assert len(images) == 1
        assert images[0].width == 100
        
        images = extract_images_from_docx(file_bytes=b"fake")
        assert len(images) == 1

@patch("app.services.extraction.image_extractor.ocr_image", return_value="ocr text")
@patch("app.services.extraction.image_extractor.Image.open")
def test_extract_images_from_pptx(mock_image_open, mock_ocr_image):
    mock_pil = MagicMock()
    mock_pil.size = (100, 100)
    mock_pil.format = "png"
    mock_image_open.return_value = mock_pil

    with patch("pptx.Presentation") as mock_pres_class:
        mock_pres = MagicMock()
        mock_pres_class.return_value = mock_pres
        
        mock_slide = MagicMock()
        mock_shape = MagicMock()
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        mock_shape.shape_type = MSO_SHAPE_TYPE.PICTURE
        mock_shape.image.blob = b"fakeimage"
        
        mock_slide.shapes = [mock_shape]
        mock_pres.slides = [mock_slide]
        
        images = extract_images_from_pptx(file_path="test.pptx")
        assert len(images) == 1
        
        images = extract_images_from_pptx(file_bytes=b"fake")
        assert len(images) == 1

@patch("app.services.extraction.image_extractor.ocr_image", return_value="ocr text")
@patch("app.services.extraction.image_extractor.Image.open")
def test_extract_images_from_file_dispatch(mock_image_open, mock_ocr_image):
    mock_pil = MagicMock()
    mock_pil.size = (100, 100)
    mock_image_open.return_value = mock_pil

    with patch("app.services.extraction.image_extractor.extract_images_from_pdf") as mock_pdf:
        extract_images_from_file(filename="test.pdf")
        mock_pdf.assert_called_once()
        
    with patch("app.services.extraction.image_extractor.extract_images_from_docx") as mock_docx:
        extract_images_from_file(filename="test.docx")
        mock_docx.assert_called_once()
        
    with patch("app.services.extraction.image_extractor.extract_images_from_pptx") as mock_pptx:
        extract_images_from_file(filename="test.pptx")
        mock_pptx.assert_called_once()

    images = extract_images_from_file(file_path="test.png")
    assert len(images) == 1
    
    images2 = extract_images_from_file(file_bytes=b"fake", filename="test.jpg")
    assert len(images2) == 1
