"""
Image extraction from documents (PDF, DOCX, PPTX).
Extracts embedded images and runs OCR on them.
"""

import io
from pathlib import Path
from typing import Optional

from PIL import Image

from app.models.schemas import ExtractedImage
from app.services.ocr_engine import ocr_image
from app.utils.logging import get_logger

logger = get_logger(__name__)


def extract_images_from_pdf(
    file_path: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
    run_ocr: bool = True,
) -> list[ExtractedImage]:
    """
    Extract all embedded images from a PDF and optionally OCR them.

    Args:
        file_path: Path to PDF.
        file_bytes: Raw PDF bytes.
        run_ocr: Whether to run OCR on extracted images.

    Returns:
        List of ExtractedImage with OCR text.
    """
    import fitz

    if file_bytes:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    elif file_path:
        doc = fitz.open(file_path)
    else:
        raise ValueError("Provide file_path or file_bytes")

    images: list[ExtractedImage] = []
    image_index = 0

    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            image_list = page.get_images(full=True)

            for img_info in image_list:
                xref = img_info[0]
                try:
                    base_image = doc.extract_image(xref)
                    if not base_image:
                        continue

                    img_bytes = base_image["image"]
                    img_ext = base_image.get("ext", "png")
                    width = base_image.get("width", 0)
                    height = base_image.get("height", 0)

                    # Skip very small images (icons, bullets, etc.)
                    if width < 50 or height < 50:
                        continue

                    ocr_text = ""
                    if run_ocr:
                        try:
                            pil_image = Image.open(io.BytesIO(img_bytes))
                            ocr_text = ocr_image(image=pil_image)
                        except Exception as e:
                            logger.warning(
                                "image_ocr_failed",
                                page=page_num + 1,
                                error=str(e),
                            )

                    images.append(
                        ExtractedImage(
                            image_index=image_index,
                            page_number=page_num + 1,
                            ocr_text=ocr_text,
                            width=width,
                            height=height,
                            format=img_ext,
                        )
                    )
                    image_index += 1

                except Exception as e:
                    logger.warning(
                        "image_extraction_failed",
                        page=page_num + 1,
                        xref=xref,
                        error=str(e),
                    )
    finally:
        doc.close()

    logger.info("pdf_images_extracted", count=len(images))
    return images


def extract_images_from_docx(
    file_path: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
    run_ocr: bool = True,
) -> list[ExtractedImage]:
    """
    Extract all embedded images from a DOCX file.

    Args:
        file_path: Path to DOCX.
        file_bytes: Raw DOCX bytes.
        run_ocr: Whether to run OCR on extracted images.

    Returns:
        List of ExtractedImage with OCR text.
    """
    from docx import Document as DocxDocument

    if file_bytes:
        doc = DocxDocument(io.BytesIO(file_bytes))
    elif file_path:
        doc = DocxDocument(file_path)
    else:
        raise ValueError("Provide file_path or file_bytes")

    images: list[ExtractedImage] = []
    image_index = 0

    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                img_data = rel.target_part.blob
                pil_image = Image.open(io.BytesIO(img_data))
                width, height = pil_image.size

                # Skip tiny images
                if width < 50 or height < 50:
                    continue

                ocr_text = ""
                if run_ocr:
                    try:
                        ocr_text = ocr_image(image=pil_image)
                    except Exception as e:
                        logger.warning("docx_image_ocr_failed", error=str(e))

                images.append(
                    ExtractedImage(
                        image_index=image_index,
                        ocr_text=ocr_text,
                        width=width,
                        height=height,
                        format=pil_image.format or "unknown",
                    )
                )
                image_index += 1

            except Exception as e:
                logger.warning("docx_image_extraction_failed", error=str(e))

    logger.info("docx_images_extracted", count=len(images))
    return images


def extract_images_from_pptx(
    file_path: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
    run_ocr: bool = True,
) -> list[ExtractedImage]:
    """
    Extract images from a PPTX file.

    Args:
        file_path: Path to PPTX.
        file_bytes: Raw PPTX bytes.
        run_ocr: Whether to run OCR.

    Returns:
        List of ExtractedImage with OCR text.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if file_bytes:
        prs = Presentation(io.BytesIO(file_bytes))
    elif file_path:
        prs = Presentation(file_path)
    else:
        raise ValueError("Provide file_path or file_bytes")

    images: list[ExtractedImage] = []
    image_index = 0

    for slide_num, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img_data = shape.image.blob
                    pil_image = Image.open(io.BytesIO(img_data))
                    width, height = pil_image.size

                    if width < 50 or height < 50:
                        continue

                    ocr_text = ""
                    if run_ocr:
                        try:
                            ocr_text = ocr_image(image=pil_image)
                        except Exception as e:
                            logger.warning("pptx_image_ocr_failed", error=str(e))

                    images.append(
                        ExtractedImage(
                            image_index=image_index,
                            page_number=slide_num,
                            ocr_text=ocr_text,
                            width=width,
                            height=height,
                            format=pil_image.format or "unknown",
                        )
                    )
                    image_index += 1

                except Exception as e:
                    logger.warning("pptx_image_extraction_failed", error=str(e))

    logger.info("pptx_images_extracted", count=len(images))
    return images


def extract_images_from_file(
    file_path: Optional[str] = None,
    file_bytes: Optional[bytes] = None,
    file_type: str = "",
    filename: str = "",
    run_ocr: bool = True,
) -> list[ExtractedImage]:
    """
    Unified image extraction entry point.

    Args:
        file_path: Path to file.
        file_bytes: Raw file bytes.
        file_type: MIME type.
        filename: Original filename.
        run_ocr: Whether to OCR extracted images.

    Returns:
        List of ExtractedImage.
    """
    ext = ""
    if filename:
        ext = Path(filename).suffix.lower().lstrip(".")
    if not ext and file_path:
        ext = Path(file_path).suffix.lower().lstrip(".")

    if ext == "pdf":
        return extract_images_from_pdf(file_path=file_path, file_bytes=file_bytes, run_ocr=run_ocr)
    elif ext in ("docx", "doc"):
        return extract_images_from_docx(file_path=file_path, file_bytes=file_bytes, run_ocr=run_ocr)
    elif ext in ("pptx", "ppt"):
        return extract_images_from_pptx(file_path=file_path, file_bytes=file_bytes, run_ocr=run_ocr)
    elif ext in ("png", "jpg", "jpeg", "tiff", "bmp"):
        # The file itself is an image — OCR it directly
        if file_bytes:
            pil_image = Image.open(io.BytesIO(file_bytes))
        elif file_path:
            pil_image = Image.open(file_path)
        else:
            return []

        width, height = pil_image.size
        ocr_text = ocr_image(image=pil_image) if run_ocr else ""
        return [
            ExtractedImage(
                image_index=0,
                ocr_text=ocr_text,
                width=width,
                height=height,
                format=ext,
            )
        ]
    else:
        return []
